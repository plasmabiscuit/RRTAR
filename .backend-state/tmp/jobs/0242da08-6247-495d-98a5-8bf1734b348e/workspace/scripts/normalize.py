from __future__ import annotations

"""
Normalize validated XML files into import_manifest.json and import_manifest.xlsx.

Each manifest entry represents one person to be added to the prime Workspace form.
Source PDPI → add_as_key_person (core rule: source PDPI is never the target PDPI).
Source KeyPerson[N] → add_as_key_person.
"""

import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
XML_DIR = BASE / "extracted-xml"
ATT_DIR = BASE / "extracted-attachments"
INPUT_DIR = BASE / "input-pdfs"
REVIEW_DIR = BASE / "review"
CONFIG_DIR = BASE / "config"
DATA_DIR = BASE / "data"

FORM_NS = "http://apply.grants.gov/forms/RR_KeyPersonExpanded_4_0-V4.0"
ATT_NS = "http://apply.grants.gov/system/Attachments-V1.0"
GLOB_LIB_NS = "http://apply.grants.gov/system/GlobalLibrary-V2.0"
GLOB_NS = "http://apply.grants.gov/system/Global-V1.0"


def sha1_b64(path: Path) -> Optional[str]:
    if not path or not path.exists():
        return None
    import base64
    digest = hashlib.sha1(path.read_bytes()).digest()
    return base64.b64encode(digest).decode()


def _t(el: etree._Element, tag: str, ns: str = FORM_NS) -> str:
    if el is None:
        return ""
    child = el.find(f"{{{ns}}}{tag}")
    return (child.text or "").strip() if child is not None else ""


def _form_ns(el: Optional[etree._Element]) -> str:
    return etree.QName(el.tag).namespace if el is not None else FORM_NS


def _resolve_attachment(
    desc_pattern: str,
    att_dir: Path,
    last_name: str = "",
    label: str = "",
) -> Tuple[Optional[str], Optional[str]]:
    """Return (path_str, sha1_b64) for an extracted attachment file.

    Tries short human names first (e.g. Zhan-Biosketch.pdf, Zhan-Biosketch_2.pdf),
    then falls back to the legacy long desc-based names for backwards compatibility.
    """
    if att_dir.is_dir():
        if last_name and label:
            # Match any file starting with "{LastName}-{Label}" (handles _2 collision suffix)
            prefix = f"{last_name}-{label}"
            matches = sorted(p for p in att_dir.iterdir()
                             if p.stem.startswith(prefix) and p.suffix == ".pdf")
            if matches:
                return str(matches[0].relative_to(BASE)), sha1_b64(matches[0])
        elif label:
            candidate = att_dir / f"{label}.pdf"
            if candidate.exists():
                return str(candidate.relative_to(BASE)), sha1_b64(candidate)

    # Legacy: desc-based long names (endswith avoids glob bracket interpretation)
    safe_base = re.sub(r"[^\w.\-\[\]]", "_", desc_pattern) + ".pdf"
    candidate = att_dir / safe_base
    if candidate.exists():
        return str(candidate.relative_to(BASE)), sha1_b64(candidate)
    if att_dir.is_dir():
        matches = sorted(p for p in att_dir.iterdir() if p.name.endswith(safe_base))
        if matches:
            return str(matches[0].relative_to(BASE)), sha1_b64(matches[0])
    return None, None


def _clean_ws(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _norm_name(text: str) -> str:
    return _clean_ws(text).lower()


def _contact_campus_box_line(contact: dict) -> str:
    raw = contact.get("campusBox")
    if raw is None:
        raw = contact.get("Campus Box")
    box = _clean_ws(str(raw or ""))
    if not box:
        return ""
    return f"Campus Box {box}"


def _pdf_text(path: Path) -> str:
    try:
        proc = subprocess.run(
            ["pdftotext", str(path), "-"],
            check=True,
            capture_output=True,
            text=True,
        )
        return proc.stdout
    except Exception:
        return ""


def _split_comma_name(raw: str) -> tuple[str, str, str]:
    raw = _clean_ws(raw)
    if not raw:
        return "", "", ""
    if "," in raw:
        last, rest = [part.strip() for part in raw.split(",", 1)]
    else:
        parts = raw.split()
        if len(parts) == 1:
            return parts[0], "", ""
        last, rest = parts[-1], " ".join(parts[:-1])
    rest_parts = rest.split()
    first = rest_parts[0] if rest_parts else ""
    middle = " ".join(rest_parts[1:]) if len(rest_parts) > 1 else ""
    return first, middle, last


def _match_line(text: str, pattern: str) -> str:
    m = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
    return _clean_ws(m.group(1)) if m else ""


def _classify_attachment_pdf(text: str) -> Optional[str]:
    upper = text.upper()
    if "BIOGRAPHICAL SKETCH" in upper:
        return "biosketch"
    if "CURRENT AND PENDING" in upper or "C&P(O)S" in upper:
        return "current_pending_support"
    return None


def _parse_biosketch_pdf(path: Path, text: str) -> dict:
    raw_name = _match_line(text, r"^\s*NAME:\s*(.+)$")
    first, middle, last = _split_comma_name(raw_name)
    title = _match_line(text, r"^\s*POSITION TITLE:\s*(.+)$")
    org = _match_line(text, r"^\s*PRIMARY ORGANIZATION AND LOCATION:\s*(.+)$")
    return {
        "kind": "biosketch",
        "source_pdf": path.stem,
        "person": {
            "first_name": first,
            "middle_name": middle,
            "last_name": last,
            "title": title,
            "organization_name": org,
        },
        "attachment": {
            "path": str(path.relative_to(BASE)),
            "sha1_base64": sha1_b64(path),
        },
    }


def _parse_support_pdf(path: Path, text: str) -> dict:
    raw_name = _match_line(text, r"^\s*\*?NAME:\s*(.+)$")
    first, middle, last = _split_comma_name(raw_name)
    title = _match_line(text, r"^\s*\*?POSITION TITLE:\s*(.+)$")
    org = _match_line(text, r"^\s*\*?ORGANIZATION AND LOCATION:\s*(.+)$")
    return {
        "kind": "current_pending_support",
        "source_pdf": path.stem,
        "person": {
            "first_name": first,
            "middle_name": middle,
            "last_name": last,
            "title": title,
            "organization_name": org,
        },
        "attachment": {
            "path": str(path.relative_to(BASE)),
            "sha1_base64": sha1_b64(path),
        },
    }


def _parse_attachment_only_pdf(path: Path) -> Optional[dict]:
    text = _pdf_text(path)
    if not text:
        return None
    kind = _classify_attachment_pdf(text)
    if kind == "biosketch":
        return _parse_biosketch_pdf(path, text)
    if kind == "current_pending_support":
        return _parse_support_pdf(path, text)
    return None


def _attachment_person_key(parsed: dict) -> Tuple[str, str]:
    person = parsed.get("person", {})
    return (
        (person.get("first_name") or "").lower().strip(),
        (person.get("last_name") or "").lower().strip(),
    )


def _load_contacts() -> Tuple[Dict[str, dict], Dict[str, dict]]:
    path = DATA_DIR / "contacts.json"
    if not path.exists():
        return {}, {}
    data = json.loads(path.read_text())
    if not isinstance(data, dict):
        return {}, {}

    by_email: dict[str, dict] = {}
    by_name: dict[str, dict] = {}
    for key, value in data.items():
        if not isinstance(value, dict):
            continue
        email = _clean_ws(str(value.get("email") or key)).lower()
        if email:
            by_email[email] = value
        names = {
            _norm_name(str(value.get("fullName") or "")),
            _norm_name(
                " ".join(
                    part for part in [
                        str(value.get("firstName") or ""),
                        str(value.get("lastName") or ""),
                    ] if _clean_ws(part)
                )
            ),
        }
        for full_name in names:
            if full_name and full_name not in by_name:
                by_name[full_name] = value
    return by_email, by_name


def _fill_if_blank(target: dict, key: str, value: str) -> None:
    if value and not _clean_ws(str(target.get(key, ""))):
        target[key] = value


def _enrich_from_contacts(entries: list[dict]) -> list[dict]:
    by_email, by_name = _load_contacts()
    if not by_email and not by_name:
        return entries

    for entry in entries:
        person = entry.get("person", {})
        email = _clean_ws(str(person.get("email") or "")).lower()
        full_name = _norm_name(
            " ".join(
                part for part in [
                    str(person.get("first_name") or ""),
                    str(person.get("middle_name") or ""),
                    str(person.get("last_name") or ""),
                ] if _clean_ws(part)
            )
        )
        first_last = _norm_name(
            " ".join(
                part for part in [
                    str(person.get("first_name") or ""),
                    str(person.get("last_name") or ""),
                ] if _clean_ws(part)
            )
        )

        contact = by_email.get(email) if email else None
        matched_by = "email" if contact else None
        if not contact and full_name:
            contact = by_name.get(full_name)
            matched_by = "full name" if contact else None
        if not contact and first_last:
            contact = by_name.get(first_last)
            matched_by = "first+last name" if contact else None
        if not contact:
            continue

        _fill_if_blank(person, "email", _clean_ws(str(contact.get("email") or "")))
        _fill_if_blank(person, "first_name", _clean_ws(str(contact.get("firstName") or "")))
        _fill_if_blank(person, "last_name", _clean_ws(str(contact.get("lastName") or "")))
        _fill_if_blank(person, "title", _clean_ws(str(contact.get("jobTitle") or "")))
        _fill_if_blank(person, "department", _clean_ws(str(contact.get("department") or "")))
        _fill_if_blank(person, "division", _clean_ws(str(contact.get("Parent") or contact.get("unit") or "")))
        _fill_if_blank(person, "phone", _clean_ws(str(contact.get("workPhone") or "")))
        _fill_if_blank(person, "organization_name", "Tennessee Technological University")
        _fill_if_blank(person.get("address", {}), "street1", "1 William L. Jones Dr")
        _fill_if_blank(person.get("address", {}), "street2", _contact_campus_box_line(contact))

        entry.setdefault("validation", {}).setdefault("warnings", []).append(
            f"Missing fields enriched from data/contacts.json matched by {matched_by}."
        )

    return entries


def _attachment_only_entries(agency: str) -> list[dict]:
    if not INPUT_DIR.exists():
        return []

    grouped: dict[tuple[str, str], dict] = {}
    for pdf_path in sorted(INPUT_DIR.glob("*.pdf")):
        parsed = _parse_attachment_only_pdf(pdf_path)
        if not parsed:
            continue

        key = _attachment_person_key(parsed)
        if not any(key):
            key = ("", pdf_path.stem.lower())

        bucket = grouped.setdefault(key, {
            "source_pdf": parsed["source_pdf"],
            "person": {
                "prefix": "",
                "first_name": "",
                "middle_name": "",
                "last_name": "",
                "suffix": "",
                "title": "",
                "organization_name": "",
                "department": "",
                "division": "",
                "credential": "",
                "project_role": "Other Professional",
                "other_project_role_category": "",
                "degree_type": "",
                "degree_year": "",
                "address": {
                    "street1": "",
                    "street2": "",
                    "city": "",
                    "county": "",
                    "state": "",
                    "province": "",
                    "country": "",
                    "postal_code": "",
                },
                "phone": "",
                "fax": "",
                "email": "",
            },
            "attachments": {
                "biosketch": {
                    "required": True,
                    "source": "none",
                    "path": None,
                    "sha1_base64": None,
                },
                "current_pending_support": {
                    "required": agency == "NIH_PHS",
                    "source": "none",
                    "path": None,
                    "sha1_base64": None,
                },
            },
            "validation": {
                "schema_valid": True,
                "business_rules_valid": True,
                "warnings": [],
            },
        })

        for field in ["first_name", "middle_name", "last_name", "title", "organization_name"]:
            if parsed["person"].get(field) and not bucket["person"].get(field):
                bucket["person"][field] = parsed["person"][field]

        att_key = parsed["kind"]
        bucket["attachments"][att_key] = {
            "required": att_key == "biosketch" or agency == "NIH_PHS",
            "source": "input-pdfs",
            "path": parsed["attachment"]["path"],
            "sha1_base64": parsed["attachment"]["sha1_base64"],
        }

    entries: list[dict] = []
    for bucket in grouped.values():
        warnings = bucket["validation"]["warnings"]
        if not bucket["attachments"]["biosketch"]["path"]:
            warnings.append("Biosketch file not found for attachment-only fallback person")
        if agency == "NIH_PHS" and not bucket["attachments"]["current_pending_support"]["path"]:
            warnings.append("Current & Pending Support file not found for attachment-only fallback person")
        warnings.append("Generated from standalone biosketch/current-pending PDF(s); review role and contact fields.")
        if not bucket["person"]["organization_name"]:
            warnings.append("Organization could not be inferred from standalone attachment PDF")
        if not bucket["person"]["title"]:
            warnings.append("Title could not be inferred from standalone attachment PDF")
        if not bucket["person"]["first_name"] or not bucket["person"]["last_name"]:
            warnings.append("Name could not be fully inferred from standalone attachment PDF")

        entries.append({
            "source_pdf": bucket["source_pdf"],
            "source_element": "AttachmentFallback",
            "source_index": 0,
            "target_action": "add_as_key_person",
            "exclude": False,
            "person": bucket["person"],
            "attachments": bucket["attachments"],
            "validation": {
                "schema_valid": True,
                "business_rules_valid": not any("not found" in w.lower() for w in warnings),
                "warnings": warnings,
            },
        })
    return entries


def _parse_person(person_el: etree._Element, source_element: str,
                  source_index: int, source_pdf: str,
                  pdf_stem: str, agency: str,
                  role_map: dict) -> dict:
    """Build a manifest entry dict from a PDPI or KeyPerson XML element."""
    form_ns = _form_ns(person_el)
    form_local = etree.QName(person_el.getparent().tag).localname if person_el.getparent() is not None else "RR_KeyPersonExpanded_4_0"
    profile = person_el.find(f"{{{form_ns}}}Profile")
    name_el = profile.find(f"{{{form_ns}}}Name") if profile is not None else None
    addr_el = profile.find(f"{{{form_ns}}}Address") if profile is not None else None

    def f(tag, ns=FORM_NS):
        effective_ns = form_ns if ns == FORM_NS else ns
        return _t(profile, tag, effective_ns) if profile is not None else ""

    def n(tag):
        return _t(name_el, tag, GLOB_LIB_NS) if name_el is not None else ""

    # Address children are in GLOB_LIB_NS; Name children also use GLOB_LIB_NS
    def a(tag):
        return _t(addr_el, tag, GLOB_LIB_NS) if addr_el is not None else ""

    raw_role = f("ProjectRole")
    other_role_cat = f("OtherProjectRoleCategory")
    role, warnings = _map_role(raw_role, other_role_cat, agency, role_map)

    # Resolve biosketch and support attachments by /Desc slot convention
    # PDPI[0].mandatoryFile0 → biosketch for PDPI
    # PDPI[0].optionalFile0  → support for PDPI
    # KeyPerson[N].optionalFile0 → biosketch for KeyPerson[N]
    # KeyPerson[N].optionalFile1 → support for KeyPerson[N]
    if source_element == "PDPI":
        bio_desc = f"{form_local}_P1.PDPI[0].mandatoryFile0"
        sup_desc = f"{form_local}_P1.PDPI[0].optionalFile0"
    else:
        bio_desc = f"{form_local}_P1.KeyPerson[{source_index}].optionalFile0"
        sup_desc = f"{form_local}_P1.KeyPerson[{source_index}].optionalFile1"

    last_name_val = n("LastName")
    bio_path, bio_sha1 = _resolve_attachment(bio_desc, ATT_DIR, last_name_val, "Biosketch")
    sup_path, sup_sha1 = _resolve_attachment(sup_desc, ATT_DIR, last_name_val, "Current_Pending")

    if not bio_path:
        # Try to match by filename from XML
        bio_el = (profile.find(f"{{{form_ns}}}BioSketchsAttached/{{{form_ns}}}BioSketchAttached")
                  if profile is not None else None)
        if bio_el is not None:
            fn = _t(bio_el, "FileName", ATT_NS)
            if fn:
                for p in ATT_DIR.glob("*.pdf"):
                    if fn in p.name:
                        bio_path = str(p.relative_to(BASE))
                        bio_sha1 = sha1_b64(p)
                        warnings.append(f"Biosketch matched by filename fallback: {p.name}")
                        break

    if bio_path is None:
        warnings.append("Biosketch file not found in extracted-attachments/ — must be supplied manually")
    if sup_path is None and agency == "NIH_PHS":
        warnings.append("Current & Pending Support file not found — required for NIH submissions")

    cred = f("Credential")
    if not cred and agency == "NIH_PHS":
        warnings.append("Credential (eRA Commons username) is blank — required for NIH submissions")

    # Agency-agnostic data quality warnings
    if not n("FirstName") or not n("LastName"):
        warnings.append("Name is incomplete — first or last name is blank")
    if not f("Email"):
        warnings.append("Email address is blank")
    if "Other" in role and not other_role_cat:
        warnings.append("Role is 'Other (Specify)' but Other Project Role Category is blank")

    return {
        "source_pdf": source_pdf,
        "source_element": source_element,
        "source_index": source_index,
        "target_action": "add_as_key_person",
        "exclude": False,
        "person": {
            "prefix": n("PrefixName"),
            "first_name": n("FirstName"),
            "middle_name": n("MiddleName"),
            "last_name": n("LastName"),
            "suffix": n("SuffixName"),
            "title": f("Title"),
            "organization_name": f("OrganizationName"),
            "department": f("Department"),
            "division": f("Division"),
            "credential": cred,
            "project_role": role,
            "other_project_role_category": other_role_cat,
            "degree_type": f("DegreeType"),
            "degree_year": f("DegreeYear"),
            "address": {
                "street1":     a("Street1"),
                "street2":     a("Street2"),
                "city":        a("City"),
                "county":      a("County"),
                "state":       a("State"),
                "province":    a("Province"),
                "country":     a("Country"),
                "postal_code": a("ZipPostalCode"),
            },
            "phone": f("Phone"),
            "fax":   f("Fax"),
            "email": f("Email"),
        },
        "attachments": {
            "biosketch": {
                "required": True,
                "source": "extracted" if bio_path else "none",
                "path": bio_path,
                "sha1_base64": bio_sha1,
            },
            "current_pending_support": {
                "required": agency == "NIH_PHS",
                "source": "extracted" if sup_path else "none",
                "path": sup_path,
                "sha1_base64": sup_sha1,
            },
        },
        "validation": {
            "schema_valid": True,
            "business_rules_valid": len([w for w in warnings if "missing" in w.lower()]) == 0,
            "warnings": warnings,
        },
    }


def _map_role(raw_role: str, other_cat: str, agency: str, role_map: dict) -> tuple[str, list[str]]:
    warnings = []
    agency_rules = role_map.get(agency, {})
    if raw_role in agency_rules:
        rule = agency_rules[raw_role]
        if isinstance(rule, dict):
            mapped = rule.get("map_to", raw_role)
            if "warning" in rule:
                warnings.append(rule["warning"])
            if "other_project_role_category" in rule:
                other_cat = rule["other_project_role_category"]
            return mapped, warnings
        return str(rule), warnings
    default_rules = role_map.get("default", {})
    raw = default_rules.get(raw_role, raw_role)
    return (str(raw) if not isinstance(raw, dict) else raw.get("map_to", raw_role)), warnings


def _diff_persons(a: dict, b: dict) -> list[str]:
    """Return labels for fields that differ between two person dicts."""
    diffs = []
    for key, label in [
        ("organization_name", "org"),
        ("project_role",      "role"),
        ("credential",        "credential"),
        ("title",             "title"),
        ("degree_type",       "degree"),
        ("phone",             "phone"),
        ("email",             "email"),
    ]:
        if (a.get(key) or "").strip() != (b.get(key) or "").strip():
            diffs.append(label)
    addr_a = a.get("address", {}) or {}
    addr_b = b.get("address", {}) or {}
    for key, label in [("street1", "street"), ("city", "city"), ("state", "state")]:
        if (addr_a.get(key) or "").strip() != (addr_b.get(key) or "").strip():
            diffs.append(label)
    return diffs


def _dedup_check(entries: list[dict]) -> list[dict]:
    """Auto-exclude duplicate persons (matched by email or first+last name).

    The first occurrence is kept included. Subsequent matches are excluded and
    flagged with a warning. If field values differ the warning lists them so the
    user can decide whether to include the entry manually.
    """
    seen_email: dict[str, int] = {}   # normalised email → first-occurrence index
    seen_name:  dict[tuple, int] = {}  # (first, last) → first-occurrence index

    for i, e in enumerate(entries):
        p = e["person"]
        email    = (p.get("email", "") or "").lower().strip()
        first    = (p.get("first_name", "") or "").lower().strip()
        last     = (p.get("last_name",  "") or "").lower().strip()
        name_key = (first, last)

        match_idx: Optional[int] = None
        match_by: Optional[str] = None

        if email and email in seen_email:
            match_idx = seen_email[email]
            match_by  = "email"
        elif first and last and name_key in seen_name:
            match_idx = seen_name[name_key]
            match_by  = "name"

        if match_idx is not None:
            prior = entries[match_idx]
            src   = Path(prior.get("source_pdf", "")).name or prior.get("source_pdf", "?")
            label = f"entry #{match_idx + 1} ({src})"
            diffs = _diff_persons(p, prior["person"])

            e["exclude"] = True
            if diffs:
                e["validation"]["warnings"].append(
                    f"Possible duplicate of {label} matched by {match_by} — "
                    f"excluded for review. Differing fields: {', '.join(diffs)}"
                )
            else:
                e["validation"]["warnings"].append(
                    f"Exact duplicate of {label} matched by {match_by} — excluded."
                )
        else:
            if email:
                seen_email[email] = i
            if first and last:
                seen_name[name_key] = i

    return entries


def _write_xlsx(entries: list[dict], out_path: Path, agency: str = "default") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Import Manifest"

    headers = [
        "Include", "Source PDF", "Source Element",
        "First Name", "Last Name", "Email", "Organization",
        "Role", "Other Role", "Credential",
        "Biosketch", "Support", "Warnings",
    ]
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row_i, e in enumerate(entries, 2):
        p = e["person"]
        include = not e.get("exclude", False)
        raw_warnings = e["validation"].get("warnings", [])
        if agency != "NIH_PHS":
            raw_warnings = [w for w in raw_warnings if "NIH" not in w and "eRA" not in w]
        warnings = "; ".join(raw_warnings)
        bio = (e["attachments"]["biosketch"].get("path") or "")
        sup = (e["attachments"]["current_pending_support"].get("path") or "")

        values = [
            include,
            Path(e["source_pdf"]).name,
            e["source_element"],
            p["first_name"],
            p["last_name"],
            p["email"],
            p["organization_name"],
            p["project_role"],
            p.get("other_project_role_category", ""),
            p["credential"],
            Path(bio).name if bio else "",
            Path(sup).name if sup else "",
            warnings,
        ]
        excluded_fill = PatternFill(start_color="FFE8E8", end_color="FFE8E8", fill_type="solid")
        warn_fill     = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            cell.alignment = Alignment(wrap_text=True)
            if col == 1:
                cell.font = Font(bold=True)
            if not include:
                cell.fill = excluded_fill
            if warnings and col == 13:
                cell.fill = warn_fill

    # Column widths
    widths = [10, 28, 16, 16, 16, 28, 28, 18, 18, 18, 30, 30, 50]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    wb.save(out_path)


def run(agency: str = "default") -> list[dict]:
    xmls = sorted(XML_DIR.glob("*.xml"))

    role_map_path = CONFIG_DIR / "role_mapping.json"
    if not role_map_path.exists():
        print(f"ERROR: {role_map_path} not found. Restore it from version control.")
        return []
    role_map = json.loads(role_map_path.read_text())
    REVIEW_DIR.mkdir(exist_ok=True)

    entries = []
    for xml_path in xmls:
        print(f"  normalizing {xml_path.name} …")
        tree = etree.parse(str(xml_path))
        root = tree.getroot()
        source_pdf = xml_path.stem  # approximate; run_log has the real path

        kp_idx = 0
        for child in root:
            local = etree.QName(child.tag).localname
            if local == "PDPI":
                entry = _parse_person(child, "PDPI", 0, source_pdf, xml_path.stem, agency, role_map)
                entries.append(entry)
                print(f"    PDPI → {entry['person']['first_name']} {entry['person']['last_name']}")
            elif local == "KeyPerson":
                entry = _parse_person(child, "KeyPerson", kp_idx, source_pdf, xml_path.stem, agency, role_map)
                entries.append(entry)
                print(f"    KeyPerson[{kp_idx}] → {entry['person']['first_name']} {entry['person']['last_name']}")
                kp_idx += 1

    fallback_entries = _attachment_only_entries(agency)
    for entry in fallback_entries:
        entries.append(entry)
        print(
            "  attachment fallback → "
            f"{entry['person']['first_name']} {entry['person']['last_name']}".strip()
        )

    if not entries:
        print(f"No XML files in {XML_DIR} and no standalone biosketch/current-pending PDFs in {INPUT_DIR}.")
        return []

    entries = _enrich_from_contacts(entries)
    entries = _dedup_check(entries)

    manifest_path = REVIEW_DIR / "import_manifest.json"
    manifest_path.write_text(json.dumps(entries, indent=2))
    print(f"\n  manifest → {manifest_path.relative_to(BASE)}")

    xlsx_path = REVIEW_DIR / "import_manifest.xlsx"
    _write_xlsx(entries, xlsx_path, agency)
    print(f"  review   → {xlsx_path.relative_to(BASE)}")
    print(f"\n{len(entries)} persons total. Review the Excel file, set Include=False to exclude, then run automate.")

    return entries


if __name__ == "__main__":
    agency = sys.argv[1] if len(sys.argv) > 1 else "default"
    run(agency)
