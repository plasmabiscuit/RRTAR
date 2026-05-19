from __future__ import annotations

"""
Normalize extracted R&R Budget XML files into budget_manifest.json and budget_manifest.xlsx.

Each manifest entry represents one source budget PDF, containing organization info and
all budget period data (Sections A–K) ready for review and webform entry.
"""

import hashlib
import importlib.util
import json
import os
import sys
from pathlib import Path
from typing import Optional, Tuple

BASE = Path(__file__).resolve().parents[1]
XML_DIR  = BASE / "extracted-budget-xml"
ATT_DIR  = BASE / "extracted-budget-attachments"
REVIEW_DIR = BASE / "review"
_VENV_PYTHON = BASE / ".venv" / "bin" / "python3"


def _ensure_modules(modules: list[str]) -> None:
    missing = [name for name in modules if importlib.util.find_spec(name) is None]
    if not missing:
        return
    if _VENV_PYTHON.exists() and sys.executable != str(_VENV_PYTHON):
        os.execv(str(_VENV_PYTHON), [str(_VENV_PYTHON), __file__, *sys.argv[1:]])
    missing_csv = ", ".join(missing)
    raise ImportError(f"Missing required module(s): {missing_csv}")


_ensure_modules(["lxml", "openpyxl"])

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORM_NS     = "http://apply.grants.gov/forms/RR_Budget_3_0-V3.0"
GLOB_LIB_NS = "http://apply.grants.gov/system/GlobalLibrary-V2.0"


def sha1_b64(path: Path) -> Optional[str]:
    if not path or not path.exists():
        return None
    import base64
    return base64.b64encode(hashlib.sha1(path.read_bytes()).digest()).decode()


def _t(el: Optional[etree._Element], tag: str, ns: str = FORM_NS) -> str:
    if el is None:
        return ""
    child = el.find(f"{{{ns}}}{tag}")
    return (child.text or "").strip() if child is not None else ""


def _parse_personnel_row(el: Optional[etree._Element]) -> dict:
    if el is None:
        return {}
    return {
        "number_of_personnel": _t(el, "NumberOfPersonnel"),
        "project_role":        _t(el, "ProjectRole"),
        "calendar_months":     _t(el, "CalendarMonths"),
        "academic_months":     _t(el, "AcademicMonths"),
        "summer_months":       _t(el, "SummerMonths"),
        "requested_salary":    _t(el, "RequestedSalary"),
        "fringe_benefits":     _t(el, "FringeBenefits"),
        "funds_requested":     _t(el, "FundsRequested"),
    }


def _is_zeroish(value: str) -> bool:
    v = (value or "").strip()
    return v in {"", "0", "0.0", "0.00"}


def _has_meaningful_personnel_row(row: dict) -> bool:
    numeric_keys = [
        "number_of_personnel",
        "calendar_months",
        "academic_months",
        "summer_months",
        "requested_salary",
        "fringe_benefits",
        "funds_requested",
    ]
    return any(not _is_zeroish(str(row.get(key, ""))) for key in numeric_keys)


def _parse_key_person(kp_el: etree._Element) -> dict:
    name_el = kp_el.find(f"{{{FORM_NS}}}Name")
    prefix = _t(name_el, "PrefixName", GLOB_LIB_NS) if name_el is not None else ""
    first = _t(name_el, "FirstName", GLOB_LIB_NS) if name_el is not None else ""
    middle = _t(name_el, "MiddleName", GLOB_LIB_NS) if name_el is not None else ""
    last  = _t(name_el, "LastName",  GLOB_LIB_NS) if name_el is not None else ""
    suffix = _t(name_el, "SuffixName", GLOB_LIB_NS) if name_el is not None else ""
    return {
        "prefix":           prefix.strip(),
        "first_name":       first.strip(),
        "middle_name":      middle.strip(),
        "last_name":        last.strip(),
        "suffix":           suffix.strip(),
        "project_role":     _t(kp_el, "ProjectRole"),
        "calendar_months":  _t(kp_el, "CalendarMonths"),
        "academic_months":  _t(kp_el, "AcademicMonths"),
        "summer_months":    _t(kp_el, "SummerMonths"),
        "base_salary":      _t(kp_el, "BaseSalary"),
        "requested_salary": _t(kp_el, "RequestedSalary"),
        "fringe_benefits":  _t(kp_el, "FringeBenefits"),
        "funds_requested":  _t(kp_el, "FundsRequested"),
    }


def _parse_section_a(year_el: etree._Element) -> dict:
    kps_el = year_el.find(f"{{{FORM_NS}}}KeyPersons")
    if kps_el is None:
        return {"key_persons": [], "total": ""}
    persons = [
        _parse_key_person(kp)
        for kp in kps_el.findall(f"{{{FORM_NS}}}KeyPerson")
    ]
    return {
        "key_persons": persons,
        "total":       _t(kps_el, "TotalFundForKeyPersons"),
    }


def _parse_section_b(year_el: etree._Element) -> dict:
    op = year_el.find(f"{{{FORM_NS}}}OtherPersonnel")
    if op is None:
        return {}

    predefined = {}
    for tag, key in [
        ("PostDocAssociates",    "post_doc_associates"),
        ("GraduateStudents",     "graduate_students"),
        ("UndergraduateStudents","undergraduate_students"),
        ("SecretarialClerical",  "secretarial_clerical"),
    ]:
        el = op.find(f"{{{FORM_NS}}}{tag}")
        if el is not None:
            row = _parse_personnel_row(el)
            if _has_meaningful_personnel_row(row):
                predefined[key] = row

    others = []
    for el in op.findall(f"{{{FORM_NS}}}Other"):
        row = _parse_personnel_row(el)
        if _has_meaningful_personnel_row(row):
            others.append(row)

    return {
        **predefined,
        "other_personnel":    others,
        "total_number":       _t(op, "OtherPersonnelTotalNumber"),
        "total_funds":        _t(op, "TotalOtherPersonnelFund"),
    }


def _parse_section_c(year_el: etree._Element) -> dict:
    eq_el = year_el.find(f"{{{FORM_NS}}}Equipment")
    if eq_el is None:
        return {"items": [], "total": ""}
    items = []
    for item_el in eq_el.findall(f"{{{FORM_NS}}}EquipmentList"):
        name = _t(item_el, "EquipmentItem")
        cost = _t(item_el, "FundsRequested")
        if name or cost:
            items.append({"item": name, "funds_requested": cost})
    return {"items": items, "total": _t(eq_el, "TotalFund")}


def _parse_section_d(year_el: etree._Element) -> dict:
    tr = year_el.find(f"{{{FORM_NS}}}Travel")
    if tr is None:
        return {}
    return {"total_travel_cost": _t(tr, "TotalTravelCost")}


def _parse_section_e(year_el: etree._Element) -> dict:
    pt = year_el.find(f"{{{FORM_NS}}}ParticipantTraineeSupportCosts")
    if pt is None:
        return {}
    other_el = pt.find(f"{{{FORM_NS}}}Other")
    return {
        "tuition_fees_health_insurance": _t(pt, "TuitionFeesHealthInsurance"),
        "stipends": _t(pt, "Stipends"),
        "travel": _t(pt, "Travel"),
        "subsistence": _t(pt, "Subsistence"),
        "other_description": _t(other_el, "Description") if other_el is not None else "",
        "other_cost": _t(other_el, "Cost") if other_el is not None else "",
        "number_of_participants": _t(pt, "NumberOfParticipantsTrainees"),
        "total_cost": _t(pt, "TotalCost"),
    }


def _parse_section_f(year_el: etree._Element) -> dict:
    odc = year_el.find(f"{{{FORM_NS}}}OtherDirectCosts")
    if odc is None:
        return {}
    other_items = []
    for n in range(1, 11):
        el = odc.find(f"{{{FORM_NS}}}OtherDirectCost{n}")
        if el is not None:
            desc = _t(el, "Description")
            cost = _t(el, "Cost")
            if desc or cost:
                other_items.append({"description": desc, "cost": cost})
    return {
        "materials_supplies": _t(odc, "MaterialsSupplies"),
        "publication_costs": _t(odc, "PublicationCosts"),
        "consultant_services": _t(odc, "ConsultantServices"),
        "adp_computer_services": _t(odc, "ADPComputerServices"),
        "subawards_consortium_contractual_costs": _t(odc, "SubawardsConsortiumContractualCosts"),
        "equipment_facility_rental_user_fees": _t(odc, "EquipmentFacilityRentalUserFees"),
        "alterations_renovations": _t(odc, "AlterationsRenovations"),
        "other_items":        other_items,
        "total":              _t(odc, "TotalOtherDirectCost"),
    }


def _parse_section_h(year_el: etree._Element) -> dict:
    ic_root = year_el.find(f"{{{FORM_NS}}}IndirectCosts")
    if ic_root is None:
        return {}
    costs = []
    for ic in ic_root.findall(f"{{{FORM_NS}}}IndirectCost"):
        cost_type = _t(ic, "CostType")
        rate      = _t(ic, "Rate")
        base      = _t(ic, "Base")
        fund      = _t(ic, "FundRequested")
        if cost_type or fund:
            costs.append({
                "cost_type":      cost_type,
                "rate":           rate,
                "base":           base,
                "fund_requested": fund,
            })
    return {
        "indirect_costs":   costs,
        "total":            _t(ic_root, "TotalIndirectCosts"),
        "cognizant_agency": _t(ic_root, "CognizantFederalAgency"),
    }


def _parse_period(year_el: etree._Element, index: int) -> dict:
    return {
        "period_index": index,
        "start_date":   _t(year_el, "BudgetPeriodStartDate"),
        "end_date":     _t(year_el, "BudgetPeriodEndDate"),
        "section_a":    _parse_section_a(year_el),
        "section_b":    _parse_section_b(year_el),
        "section_c":    _parse_section_c(year_el),
        "section_d":    _parse_section_d(year_el),
        "section_e":    _parse_section_e(year_el),
        "section_f":    _parse_section_f(year_el),
        "section_h":    _parse_section_h(year_el),
        "total_compensation": _t(year_el, "TotalCompensation"),
        "direct_costs":       _t(year_el, "DirectCosts"),
        "total_costs":        _t(year_el, "TotalCosts"),
        "total_costs_fee":    _t(year_el, "TotalCostsFee"),
    }


def _resolve_justification(stem: str) -> Tuple[Optional[str], Optional[str]]:
    if ATT_DIR.is_dir():
        for p in sorted(ATT_DIR.iterdir()):
            if p.stem.startswith(stem) and "BudgetJustification" in p.stem:
                return str(p.relative_to(BASE)), sha1_b64(p)
    return None, None


def _parse_budget(xml_path: Path) -> dict:
    tree = etree.parse(str(xml_path))
    root = tree.getroot()

    uei      = _t(root, "SAMUEI")
    org_name = _t(root, "OrganizationName")
    btype    = _t(root, "BudgetType")

    years = root.findall(f"{{{FORM_NS}}}BudgetYear")
    periods = [_parse_period(y, i + 1) for i, y in enumerate(years)]

    just_path, just_sha1 = _resolve_justification(xml_path.stem)

    return {
        "source_pdf":  xml_path.stem,
        "organization": {
            "uei":         uei,
            "name":        org_name,
            "budget_type": btype,
        },
        "budget_justification": {
            "path":       just_path,
            "sha1_base64": just_sha1,
        },
        "periods": periods,
    }


def _write_xlsx(entries: list[dict], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget Manifest"

    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    headers = [
        "Source PDF", "Organization", "Budget Type", "UEI",
        "Period", "Start Date", "End Date",
        "Key Persons", "Total Personnel ($)",
        "Direct Costs ($)", "Indirect Costs ($)", "Total Costs ($)",
        "Justification File",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_i = 2
    for entry in entries:
        org  = entry.get("organization", {})
        just = entry.get("budget_justification", {})
        just_name = Path(just.get("path") or "").name if just.get("path") else ""
        for period in entry.get("periods", []):
            sec_a = period.get("section_a", {})
            kp_count = len(sec_a.get("key_persons", []))
            values = [
                entry.get("source_pdf", ""),
                org.get("name", ""),
                org.get("budget_type", ""),
                org.get("uei", ""),
                period.get("period_index", ""),
                period.get("start_date", ""),
                period.get("end_date", ""),
                kp_count,
                period.get("total_compensation", ""),
                period.get("direct_costs", ""),
                period.get("section_h", {}).get("total", ""),
                period.get("total_costs", ""),
                just_name,
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_i, column=col, value=val)
                cell.alignment = Alignment(wrap_text=True)
            row_i += 1

    widths = [28, 40, 20, 16, 8, 14, 14, 12, 18, 18, 18, 18, 35]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    wb.save(out_path)


def run() -> list[dict]:
    xmls = sorted(XML_DIR.glob("*.xml"))
    if not xmls:
        print(f"No XML files in {XML_DIR}. Run extract-budget first.")
        return []

    REVIEW_DIR.mkdir(exist_ok=True)
    entries = []
    for xml_path in xmls:
        print(f"  normalizing {xml_path.name} …")
        try:
            entry = _parse_budget(xml_path)
            entries.append(entry)
            org = entry["organization"]["name"] or xml_path.stem
            periods = len(entry["periods"])
            total_kp = sum(
                len(p["section_a"]["key_persons"]) for p in entry["periods"]
            )
            print(f"    {org}  {periods} period(s)  {total_kp} key person(s)")
        except Exception as exc:
            print(f"  ERROR {xml_path.name}: {exc}")

    manifest_path = REVIEW_DIR / "budget_manifest.json"
    manifest_path.write_text(json.dumps(entries, indent=2))
    print(f"\n  manifest → {manifest_path.relative_to(BASE)}")

    xlsx_path = REVIEW_DIR / "budget_manifest.xlsx"
    _write_xlsx(entries, xlsx_path)
    print(f"  review   → {xlsx_path.relative_to(BASE)}")
    print(f"\n{len(entries)} budget source(s). Review the manifest, then run automate-budget.")
    return entries


if __name__ == "__main__":
    run()
