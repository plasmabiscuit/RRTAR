from __future__ import annotations

"""
Normalize extracted Performance Site XML files into performance_site_manifest.json
and performance_site_manifest.xlsx.
"""

import hashlib
import json
from pathlib import Path
from typing import Optional, Tuple

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
XML_DIR = BASE / "extracted-performance-site-xml"
ATT_DIR = BASE / "extracted-performance-site-attachments"
REVIEW_DIR = BASE / "review"

FORM_NS = "http://apply.grants.gov/forms/PerformanceSite_4_0-V4.0"
GLOB_LIB_NS = "http://apply.grants.gov/system/GlobalLibrary-V2.0"


def sha1_b64(path: Path) -> Optional[str]:
    if not path or not path.exists():
        return None
    import base64
    return base64.b64encode(hashlib.sha1(path.read_bytes()).digest()).decode()


def _t(el: Optional[etree._Element], tag: str, ns: str = FORM_NS) -> str:
    if el is None:
        return ""
    child = el.find(f"{{{ns}}}{tag}")
    return (child.text or "").strip() if child is not None else ""


def _addr(el: Optional[etree._Element], tag: str) -> str:
    return _t(el, tag, GLOB_LIB_NS)


def _flag(text: str) -> bool:
    return text.strip().startswith("Y:")


def _parse_site(site_el: etree._Element) -> dict:
    addr_el = site_el.find(f"{{{FORM_NS}}}Address")
    return {
        "individual": _flag(_t(site_el, "Individual")),
        "organization_name": _t(site_el, "OrganizationName"),
        "uei": _t(site_el, "SAMUEI"),
        "street1": _addr(addr_el, "Street1"),
        "street2": _addr(addr_el, "Street2"),
        "city": _addr(addr_el, "City"),
        "county": _addr(addr_el, "County"),
        "state": _addr(addr_el, "State"),
        "province": _addr(addr_el, "Province"),
        "country": _addr(addr_el, "Country"),
        "zip_postal_code": _addr(addr_el, "ZipPostalCode"),
        "congressional_district": _t(site_el, "CongressionalDistrictProgramProject"),
    }


def _resolve_attachment(stem: str) -> Tuple[Optional[str], Optional[str]]:
    if ATT_DIR.is_dir():
        for p in sorted(ATT_DIR.iterdir()):
            if p.stem.startswith(stem):
                return str(p.relative_to(BASE)), sha1_b64(p)
    return None, None


def _parse_file(xml_path: Path) -> dict:
    tree = etree.parse(str(xml_path))
    root = tree.getroot()
    primary = root.find(f"{{{FORM_NS}}}PrimarySite")
    others = root.findall(f"{{{FORM_NS}}}OtherSite")
    attachment_path, attachment_sha1 = _resolve_attachment(xml_path.stem)
    return {
        "source_pdf": xml_path.stem,
        "primary_site": _parse_site(primary) if primary is not None else {},
        "other_sites": [_parse_site(site) for site in others],
        "additional_sites_attachment": {
            "path": attachment_path,
            "sha1_base64": attachment_sha1,
        },
    }


def _write_xlsx(entries: list[dict], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Performance Site Manifest"

    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    headers = [
        "Source PDF", "Site Type", "Individual", "Organization Name", "UEI",
        "Street1", "City", "State", "Country", "ZIP / Postal Code", "Congressional District",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_i = 2
    for entry in entries:
        rows = [("Primary", entry.get("primary_site", {}))]
        rows.extend((f"Other {idx}", site) for idx, site in enumerate(entry.get("other_sites", []), 1))
        for site_type, site in rows:
            values = [
                entry.get("source_pdf", ""),
                site_type,
                "Yes" if site.get("individual") else "No",
                site.get("organization_name", ""),
                site.get("uei", ""),
                site.get("street1", ""),
                site.get("city", ""),
                site.get("state", "") or site.get("province", ""),
                site.get("country", ""),
                site.get("zip_postal_code", ""),
                site.get("congressional_district", ""),
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_i, column=col, value=val)
                cell.alignment = Alignment(wrap_text=True)
            row_i += 1

    widths = [30, 12, 10, 36, 16, 28, 18, 18, 22, 18, 18]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    wb.save(out_path)


def run() -> list[dict]:
    xmls = sorted(XML_DIR.glob("*.xml"))
    if not xmls:
        print(f"No XML files in {XML_DIR}. Run extract-performance-site first.")
        return []

    REVIEW_DIR.mkdir(exist_ok=True)
    entries = []
    for xml_path in xmls:
        print(f"  normalizing {xml_path.name} …")
        try:
            entry = _parse_file(xml_path)
            entries.append(entry)
            print(f"    other site count: {len(entry['other_sites'])}")
        except Exception as exc:
            print(f"  ERROR {xml_path.name}: {exc}")

    manifest_path = REVIEW_DIR / "performance_site_manifest.json"
    manifest_path.write_text(json.dumps(entries, indent=2))
    print(f"\n  manifest → {manifest_path.relative_to(BASE)}")

    xlsx_path = REVIEW_DIR / "performance_site_manifest.xlsx"
    _write_xlsx(entries, xlsx_path)
    print(f"  review   → {xlsx_path.relative_to(BASE)}")
    print(f"\n{len(entries)} performance site source(s). Review the manifest, then run automate-performance-site.")
    return entries


if __name__ == "__main__":
    run()
